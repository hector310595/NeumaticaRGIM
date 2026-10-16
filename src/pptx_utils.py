"""Utilidades para construir las presentaciones sobre la plantilla PUCP (python-pptx)."""
import copy
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.join(os.path.dirname(__file__), "..")
FIG = os.path.join(ROOT, "figuras")
TEMPLATE = os.path.join(ROOT, "plantilla", "Plantilla_PUCP.pptx")

NAVY = RGBColor(0x04, 0x22, 0x54)
GREEN = RGBColor(0x16, 0xC7, 0x8E)
DGREEN = RGBColor(0x0B, 0x7A, 0x53)
ORANGE = RGBColor(0xB5, 0x65, 0x1D)
RED = RGBColor(0xC0, 0x39, 0x2B)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x5F, 0x6D)
TINT = RGBColor(0xEC, 0xF0, 0xF6)
GTINT = RGBColor(0xE3, 0xF6, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

L_BLANK, L_TITLE, L_TITLE_ONLY = 0, 1, 6

# zona útil de contenido (pulgadas) entre el título y la banda de pie de página
X0, X1 = 0.6, 12.75
Y0, Y1 = 1.45, 6.05


class Deck:
    def __init__(self, footer):
        self.prs = Presentation(TEMPLATE)
        # eliminar diapositivas de ejemplo de la plantilla
        sld_ids = self.prs.slides._sldIdLst
        for sid in list(sld_ids):
            rid = sid.get(qn("r:id"))
            self.prs.part.drop_rel(rid)
            sld_ids.remove(sid)
        # texto del pie de página en los diseños
        for layout in self.prs.slide_layouts:
            for t in layout._element.iter(qn("a:t")):
                if t.text and "tulo de la presentaci" in t.text:
                    t.text = footer
        self.n = 0

    # ------------------------------------------------------------ diapositivas
    def cover(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[L_BLANK])

    def title_slide(self, title, subtitle, notes=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[L_TITLE])
        s.shapes.title.text = title
        for r in s.shapes.title.text_frame.paragraphs[0].runs:
            r.font.color.rgb = NAVY
            r.font.bold = True
        ph = s.placeholders[1]
        ph.text = subtitle
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = MUTED
                r.font.size = Pt(20)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def section(self, num, title, subtitle, notes=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[L_TITLE])
        t = s.shapes.title
        t.text_frame.text = ""
        p = t.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = "%s  " % num
        r.font.color.rgb = GREEN
        r.font.bold = True
        r = p.add_run()
        r.text = title
        r.font.color.rgb = NAVY
        r.font.bold = True
        ph = s.placeholders[1]
        ph.text = subtitle
        for pp in ph.text_frame.paragraphs:
            for rr in pp.runs:
                rr.font.color.rgb = MUTED
                rr.font.size = Pt(20)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def content(self, title, notes=None, hidden=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[L_TITLE_ONLY])
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.0)
        t.text = title
        for r in t.text_frame.paragraphs[0].runs:
            r.font.color.rgb = NAVY
            r.font.bold = True
            r.font.size = Pt(30)
        _slide_number(s)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if hidden and not os.environ.get("SHOW_HIDDEN"):
            s._element.set("show", "0")
        return s

    def save(self, path):
        self.prs.save(path)


def _slide_number(slide):
    """Copia el marcador de número de diapositiva del diseño (círculo verde del pie)."""
    from pptx.enum.shapes import PP_PLACEHOLDER
    import uuid
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes.clone_placeholder(ph)
            shp = [x for x in slide.placeholders if x.placeholder_format.idx == ph.placeholder_format.idx][0]
            p = shp.text_frame.paragraphs[0]._p
            fld = etree.SubElement(p, qn("a:fld"))
            fld.set("id", "{%s}" % str(uuid.uuid4()).upper())
            fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:rPr")).set("lang", "es-PE")
            etree.SubElement(fld, qn("a:t")).text = str(len(slide.part.package.presentation_part.presentation.slides))
            # el campo debe ir antes de endParaRPr
            end = p.find(qn("a:endParaRPr"))
            if end is not None:
                p.remove(end)
                p.append(end)


# ---------------------------------------------------------------- elementos
def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6, name=None, font=None):
    """paras: lista de str o de listas de (texto, {opciones}) para runs con formato."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        lvl = 0
        for txt, opt in runs:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
            if font or opt.get("font"):
                f.name = opt.get("font", font)
            lvl = opt.get("lvl", lvl)
        if bullets:
            _bullet(p, lvl)
    return tb


def _bullet(p, lvl=0):
    pPr = p._p.get_or_add_pPr()
    ind = 0.28 + 0.3 * lvl
    pPr.set("marL", str(int(Inches(ind))))
    pPr.set("indent", str(int(-Inches(0.24))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    c = etree.SubElement(buClr, qn("a:srgbClr"))
    c.set("val", "16C78E" if lvl == 0 else "042254")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•" if lvl == 0 else "–")


def box(slide, x, y, w, h, fill=TINT, line=None, radius=True, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def card(slide, x, y, w, h, head, body, fill=TINT, head_color=NAVY, size=16, head_size=18,
         bullets=True):
    box(slide, x, y, w, h, fill=fill)
    text(slide, x + 0.2, y + 0.15, w - 0.4, 0.45, [head], size=head_size, bold=True, color=head_color)
    if body:
        text(slide, x + 0.2, y + 0.7, w - 0.4, h - 0.8, body, size=size, bullets=bullets,
             space_after=8)


def circle_num(slide, x, y, d, label, fill=GREEN, color=WHITE, size=16):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    return shp


def image(slide, name, x, y, w, h, align="center", name_obj=None):
    """Inserta figuras/<name>.png ajustada (sin deformar) dentro de la caja."""
    path = os.path.join(FIG, name + ".png")
    iw, ih = Image.open(path).size
    r = min(w / iw, h / ih)
    W, H = iw * r, ih * r
    if align == "center":
        xx = x + (w - W) / 2
    elif align == "left":
        xx = x
    else:
        xx = x + w - W
    yy = y + (h - H) / 2
    pic = slide.shapes.add_picture(path, Inches(xx), Inches(yy), Inches(W), Inches(H))
    pic.name = name_obj or name
    return pic, (xx, yy, W, H)


def caption(slide, x, y, w, s, size=11, color=MUTED, align=PP_ALIGN.CENTER):
    return text(slide, x, y, w, 0.35, [s], size=size, color=color, align=align)


def table(slide, x, y, w, rows, col_w, size=13, head_fill=NAVY, zebra=TINT, row_h=0.4):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.size = Pt(size)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = head_fill
            else:
                r.font.color.rgb = INK
                cell.fill.solid()
                cell.fill.fore_color.rgb = zebra if i % 2 == 0 else WHITE
    return gt
