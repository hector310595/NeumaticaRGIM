"""Utilidades python-docx con el estilo del curso (A4, Arial, azul PUCP)."""
import os

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.join(os.path.dirname(__file__), "..")
FIG = os.path.join(ROOT, "figuras")
NAVY = RGBColor(0x04, 0x22, 0x54)
GREEN = RGBColor(0x0B, 0x7A, 0x53)
MUTED = RGBColor(0x55, 0x5F, 0x6D)
RED = RGBColor(0xC0, 0x39, 0x2B)


def new_doc(title, subtitle, kicker="REPRESENTACIÓN GRÁFICA PARA INGENIERÍA MECATRÓNICA"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
        h = doc.styles["Heading %d" % lvl]
        h.font.name = "Arial"
        h.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        h.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = NAVY if lvl < 3 else GREEN
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        h.paragraph_format.space_after = Pt(4)
    # pie de página con número
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("PUCP · RGIM · Representación neumática    ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    _page_field(fp)
    # encabezado de portada
    p = doc.add_paragraph()
    r = p.add_run("PONTIFICIA UNIVERSIDAD CATÓLICA DEL PERÚ")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    p = doc.add_paragraph()
    r = p.add_run(kicker)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(12)
    r.font.color.rgb = GREEN
    rule(doc)
    return doc


def _page_field(p):
    r = p.add_run()
    r.font.size = Pt(8)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._r.append(e)


def rule(doc, color="16C78E"):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "12")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)
    return p


def para(doc, runs, size=None, align=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if isinstance(runs, str):
        runs = [(runs, {})]
    for txt, o in runs:
        r = p.add_run(txt)
        r.bold = o.get("bold", False)
        r.italic = o.get("italic", False)
        if "color" in o:
            r.font.color.rgb = o["color"]
        if size or "size" in o:
            r.font.size = Pt(o.get("size", size))
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(doc, items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        runs = it if isinstance(it, list) else [(it, {})]
        for txt, o in runs:
            r = p.add_run(txt)
            r.bold = o.get("bold", False)
            if "color" in o:
                r.font.color.rgb = o["color"]
        p.paragraph_format.space_after = Pt(2)


def numbered(doc, items):
    """Lista numerada que reinicia en 1 (numeración manual con sangría francesa)."""
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("%d.\t" % i)
        r.bold = True
        r.font.color.rgb = GREEN
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        p.add_run(it)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def table(doc, rows, widths_cm, header=True, size=9.5, zebra="ECF0F6"):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, w in enumerate(widths_cm):
        t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths_cm[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            r.font.size = Pt(size)
            if header and i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, "042254")
            elif zebra and i % 2 == 0:
                shade(c, zebra)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, title, body, fill="E3F6EF", color=GREEN):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = color
    r = p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, name, width_cm, cap=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(FIG, name + ".png"), width=Cm(width_cm))
    p.paragraph_format.space_after = Pt(2)
    if cap:
        para(doc, [(cap, {"italic": True, "color": MUTED})], size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def answer_box(doc, label, height_lines=6):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    p = c.paragraphs[0]
    r = p.add_run(label)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    for _ in range(height_lines):
        c.add_paragraph()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def landscape(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    return sec


def portrait(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    return sec


def image_table(doc, header, rows, widths_cm, img_w_cm=3.0, size=9):
    """Tabla cuya primera columna es una figura (nombre en figuras/) y el resto texto.
    Cada celda de texto puede ser str o lista de (texto, {bold, color})."""
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, w in enumerate(widths_cm):
        t.columns[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.width = Cm(widths_cm[j])
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "042254")
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths_cm[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                path = os.path.join(FIG, val + ".png")
                from PIL import Image
                native_cm = Image.open(path).size[0] / 220 * 2.54  # figuras a 220 dpi
                p.add_run().add_picture(path, width=Cm(min(img_w_cm, native_cm * 0.62)))
                continue
            runs = val if isinstance(val, list) else [(val, {})]
            for txt, o in runs:
                r = p.add_run(txt)
                r.font.size = Pt(size)
                r.bold = o.get("bold", False)
                if "color" in o:
                    r.font.color.rgb = o["color"]
    # que no se parta una fila entre páginas
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit")
        trPr.append(e)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
